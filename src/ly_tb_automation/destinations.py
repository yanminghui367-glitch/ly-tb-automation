"""Destination validation, persistence, and spreadsheet import services."""

from __future__ import annotations

import io
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, BinaryIO

from .database import transaction

IMPORT_HEADERS = (
    "country_code",
    "country_name_zh",
    "country_name_en",
    "city_code",
    "city_name_zh",
    "city_name_en",
    "city_slug",
    "is_core_city",
    "priority",
    "is_enabled",
)
MAX_IMPORT_ROWS = 10_000
CODE_PATTERN = re.compile(r"^[A-Z][A-Z0-9-]{1,39}$")
SLUG_PATTERN = re.compile(r"^[a-z0-9]+(?:-[a-z0-9]+)*$")
TRUE_VALUES = {"1", "true", "yes", "y", "是", "启用"}
FALSE_VALUES = {"0", "false", "no", "n", "否", "禁用"}


@dataclass(frozen=True, slots=True)
class ImportErrorDetail:
    row_number: int
    field: str
    error_code: str
    message_zh: str
    raw_value: str | None = None


@dataclass(frozen=True, slots=True)
class DestinationInput:
    country_code: str
    country_name_zh: str
    country_name_en: str
    city_code: str
    city_name_zh: str
    city_name_en: str
    city_slug: str
    is_core_city: bool
    priority: int
    is_enabled: bool


@dataclass(slots=True)
class ImportReport:
    source_filename: str
    total_rows: int = 0
    inserted_rows: int = 0
    updated_rows: int = 0
    skipped_rows: int = 0
    errors: list[ImportErrorDetail] = field(default_factory=list)
    import_id: int | None = None

    @property
    def error_count(self) -> int:
        return len(self.errors)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _required(row_number: int, field_name: str, value: Any) -> tuple[str, ImportErrorDetail | None]:
    normalized = _text(value)
    if normalized:
        return normalized, None
    return "", ImportErrorDetail(
        row_number, field_name, "REQUIRED_FIELD_MISSING", f"{field_name} 不能为空"
    )


def _boolean(row_number: int, field_name: str, value: Any) -> tuple[bool, ImportErrorDetail | None]:
    normalized = _text(value).lower()
    if normalized in TRUE_VALUES:
        return True, None
    if normalized in FALSE_VALUES:
        return False, None
    return False, ImportErrorDetail(
        row_number,
        field_name,
        "INVALID_BOOLEAN",
        f"{field_name} 必须是 1/0、是/否或 true/false",
        _text(value),
    )


def _priority(row_number: int, value: Any) -> tuple[int, ImportErrorDetail | None]:
    try:
        priority = int(_text(value))
    except ValueError:
        priority = -1
    if 0 <= priority <= 9999:
        return priority, None
    return 100, ImportErrorDetail(
        row_number, "priority", "INVALID_PRIORITY", "priority 必须是 0 到 9999 的整数", _text(value)
    )


def validate_destination_row(
    row_number: int, raw: dict[str, Any]
) -> tuple[DestinationInput | None, list[ImportErrorDetail]]:
    """Normalize and validate one spreadsheet row without touching the database."""

    errors: list[ImportErrorDetail] = []
    values: dict[str, str] = {}
    for name in IMPORT_HEADERS[:7]:
        values[name], error = _required(row_number, name, raw.get(name))
        if error:
            errors.append(error)

    country_code = values["country_code"].upper()
    city_code = values["city_code"].upper()
    city_slug = values["city_slug"].lower()
    if country_code and not re.fullmatch(r"[A-Z]{2}", country_code):
        errors.append(
            ImportErrorDetail(
                row_number,
                "country_code",
                "INVALID_COUNTRY_CODE",
                "country_code 必须是两个英文字母",
                country_code,
            )
        )
    if city_code and not CODE_PATTERN.fullmatch(city_code):
        errors.append(
            ImportErrorDetail(
                row_number,
                "city_code",
                "INVALID_CITY_CODE",
                "city_code 只能包含大写字母、数字和连字符",
                city_code,
            )
        )
    if country_code and city_code and not city_code.startswith(f"{country_code}-"):
        errors.append(
            ImportErrorDetail(
                row_number,
                "city_code",
                "CITY_COUNTRY_MISMATCH",
                "city_code 必须以 country_code 和连字符开头",
                city_code,
            )
        )
    if city_slug and not SLUG_PATTERN.fullmatch(city_slug):
        errors.append(
            ImportErrorDetail(
                row_number,
                "city_slug",
                "INVALID_CITY_SLUG",
                "city_slug 必须是小写英文、数字或连字符",
                city_slug,
            )
        )

    is_core_city, boolean_error = _boolean(row_number, "is_core_city", raw.get("is_core_city"))
    is_enabled, enabled_error = _boolean(row_number, "is_enabled", raw.get("is_enabled"))
    priority, priority_error = _priority(row_number, raw.get("priority"))
    errors.extend(error for error in (boolean_error, enabled_error, priority_error) if error)
    if errors:
        return None, errors
    return (
        DestinationInput(
            country_code=country_code,
            country_name_zh=values["country_name_zh"],
            country_name_en=values["country_name_en"],
            city_code=city_code,
            city_name_zh=values["city_name_zh"],
            city_name_en=values["city_name_en"],
            city_slug=city_slug,
            is_core_city=is_core_city,
            priority=priority,
            is_enabled=is_enabled,
        ),
        [],
    )


def import_destination_rows(
    database_path: Path, source_filename: str, rows: list[dict[str, Any]]
) -> ImportReport:
    """Validate rows, skip invalid input, and atomically upsert valid destinations."""

    report = ImportReport(source_filename=source_filename, total_rows=len(rows))
    valid: list[tuple[int, DestinationInput]] = []
    seen_city_codes: set[str] = set()
    seen_countries: dict[str, tuple[str, str]] = {}
    for row_number, raw in enumerate(rows, start=2):
        item, errors = validate_destination_row(row_number, raw)
        if item and item.city_code in seen_city_codes:
            errors.append(
                ImportErrorDetail(
                    row_number,
                    "city_code",
                    "DUPLICATE_CITY_IN_FILE",
                    "同一文件中 city_code 重复",
                    item.city_code,
                )
            )
            item = None
        if item:
            country_names = (item.country_name_zh, item.country_name_en)
            previous_names = seen_countries.get(item.country_code)
            if previous_names is not None and previous_names != country_names:
                errors.append(
                    ImportErrorDetail(
                        row_number,
                        "country_code",
                        "COUNTRY_DATA_CONFLICT",
                        "同一文件中的国家中英文名称不一致",
                        item.country_code,
                    )
                )
                item = None
        if item:
            seen_countries[item.country_code] = (item.country_name_zh, item.country_name_en)
            seen_city_codes.add(item.city_code)
            valid.append((row_number, item))
        else:
            report.errors.extend(errors)
            report.skipped_rows += 1

    timestamp = datetime.now(UTC).isoformat()
    with transaction(database_path) as connection:
        for row_number, item in valid:
            connection.execute("SAVEPOINT destination_row")
            try:
                existing = connection.execute(
                    "SELECT id FROM cities WHERE city_code = ?", (item.city_code,)
                ).fetchone()
                connection.execute(
                    """
                    INSERT INTO countries (
                        country_code, name_zh, name_en, priority, is_enabled, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(country_code) DO UPDATE SET
                        name_zh = excluded.name_zh,
                        name_en = excluded.name_en,
                        updated_at = excluded.updated_at
                    """,
                    (
                        item.country_code,
                        item.country_name_zh,
                        item.country_name_en,
                        item.priority,
                        int(item.is_enabled),
                        timestamp,
                        timestamp,
                    ),
                )
                country_id = connection.execute(
                    "SELECT id FROM countries WHERE country_code = ?", (item.country_code,)
                ).fetchone()["id"]
                connection.execute(
                    """
                    INSERT INTO cities (
                        city_code, country_id, name_zh, name_en, slug, is_core_city,
                        priority, is_enabled, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(city_code) DO UPDATE SET
                        country_id = excluded.country_id,
                        name_zh = excluded.name_zh,
                        name_en = excluded.name_en,
                        slug = excluded.slug,
                        is_core_city = excluded.is_core_city,
                        priority = excluded.priority,
                        is_enabled = excluded.is_enabled,
                        updated_at = excluded.updated_at
                    """,
                    (
                        item.city_code,
                        country_id,
                        item.city_name_zh,
                        item.city_name_en,
                        item.city_slug,
                        int(item.is_core_city),
                        item.priority,
                        int(item.is_enabled),
                        timestamp,
                        timestamp,
                    ),
                )
                if existing:
                    report.updated_rows += 1
                else:
                    report.inserted_rows += 1
                connection.execute("RELEASE SAVEPOINT destination_row")
            except sqlite3.IntegrityError as exc:
                connection.execute("ROLLBACK TO SAVEPOINT destination_row")
                connection.execute("RELEASE SAVEPOINT destination_row")
                report.skipped_rows += 1
                report.errors.append(
                    ImportErrorDetail(
                        row_number,
                        "city_slug",
                        "DESTINATION_CONFLICT",
                        "目的地与已有国家或城市数据冲突",
                        str(exc),
                    )
                )

        cursor = connection.execute(
            """
            INSERT INTO destination_imports (
                source_filename, total_rows, inserted_rows, updated_rows,
                skipped_rows, error_count, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                report.source_filename,
                report.total_rows,
                report.inserted_rows,
                report.updated_rows,
                report.skipped_rows,
                report.error_count,
                timestamp,
            ),
        )
        report.import_id = cursor.lastrowid
        connection.executemany(
            """
            INSERT INTO destination_import_errors (
                import_id, row_number, field, error_code, message_zh, raw_value, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    report.import_id,
                    error.row_number,
                    error.field,
                    error.error_code,
                    error.message_zh,
                    error.raw_value,
                    timestamp,
                )
                for error in report.errors
            ],
        )
    return report


def list_destinations(database_path: Path) -> list[sqlite3.Row]:
    with transaction(database_path) as connection:
        return list(
            connection.execute(
                """
                SELECT
                    countries.country_code,
                    countries.name_zh AS country_name_zh,
                    countries.name_en AS country_name_en,
                    cities.city_code,
                    cities.name_zh AS city_name_zh,
                    cities.name_en AS city_name_en,
                    cities.slug,
                    cities.is_core_city,
                    cities.priority,
                    cities.image_status,
                    cities.is_enabled
                FROM cities
                JOIN countries ON countries.id = cities.country_id
                ORDER BY countries.priority, cities.priority, cities.name_zh
                """
            )
        )


def destination_counts(database_path: Path) -> tuple[int, int]:
    with transaction(database_path) as connection:
        countries = connection.execute("SELECT COUNT(*) FROM countries").fetchone()[0]
        cities = connection.execute("SELECT COUNT(*) FROM cities").fetchone()[0]
    return countries, cities


def read_destination_workbook(file: BinaryIO) -> list[dict[str, Any]]:
    """Read the first worksheet and return dictionaries using the required headers."""

    from openpyxl import load_workbook

    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise ValueError("Excel 不包含可读取的工作表")

        iterator = worksheet.iter_rows(values_only=True)
        headers = tuple(_text(value) for value in next(iterator, ()))
        missing = [name for name in IMPORT_HEADERS if name not in headers]
        if missing:
            raise ValueError(f"Excel 缺少字段: {', '.join(missing)}")
        indexes = {name: headers.index(name) for name in IMPORT_HEADERS}
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if not any(value is not None and _text(value) for value in values):
                continue
            rows.append(
                {
                    name: values[index] if index < len(values) else None
                    for name, index in indexes.items()
                }
            )
            if len(rows) > MAX_IMPORT_ROWS:
                raise ValueError(f"单次导入不能超过 {MAX_IMPORT_ROWS} 行")
        return rows
    finally:
        workbook.close()


def build_destination_template() -> bytes:
    """Build a Chinese-friendly Excel import template with one example row."""

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    try:
        worksheet = workbook.active
        if worksheet is None:
            raise RuntimeError("无法创建 Excel 工作表")

        worksheet.title = "目的地导入"
        worksheet.append(IMPORT_HEADERS)
        worksheet.append(("JP", "日本", "Japan", "JP-TOKYO", "东京", "Tokyo", "tokyo", 1, 10, 1))
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:J{worksheet.max_row}"
        for column_index in range(1, len(IMPORT_HEADERS) + 1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = 20
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
